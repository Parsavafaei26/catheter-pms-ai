"""Build the blind manual-coding workbook for step 5 (validation).

Picks 100 of the 300 AI-coded reports (stratified by year, seed 5, voided duplicates excluded)
and writes them to data/step5_manual_coding.xlsx WITHOUT the AI labels.
"""
import csv, random, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
P = ROOT / "data" / "processed"
sys.path.insert(0, str(ROOT / "scripts"))
from prepare_batches import dedupe, mfr_summary  # noqa: E402

ai = {r["report_number"]: r for r in csv.DictReader(open(P / "ai_labels_full.csv", encoding="utf-8"))}
rep = {r["report_number"]: r for r in csv.DictReader(open(P / "maude_reports.csv", encoding="utf-8"))}
pool = [k for k, v in ai.items() if "VOIDED" not in v["note"]]
random.seed(5)
pick = []
for y, n in (("2023", 33), ("2024", 33), ("2025", 34)):
    ids = sorted(k for k in pool if rep[k]["date_received"].startswith(y))
    pick += random.sample(ids, n)
random.shuffle(pick)
(P / "validation_ids.txt").write_text("\n".join(pick) + "\n")

wb = Workbook()
ins = wb.active
ins.title = "Instructions"
lines = [
    ("Step 5: blind manual coding", True),
    ("Code each report yourself using docs/02_failure_mode_categories.md (v1.1).", False),
    ("Do NOT open data/processed/ai_labels_full.csv until all 100 are done.", True),
    ("Use only the report text, as the AI did. FDA problem codes are hints only.", False),
    ("", False),
    ("Primary / secondary failure mode", True),
    ("FM01 Catheter body break (tube inside or outside body, incl. tube-hub joint)", False),
    ("FM02 External component (hub, extension, clamp, connector, valve, septum, wings)", False),
    ("FM03 Occlusion / failure to infuse or aspirate", False),
    ("FM04 Kink, deformation, ballooning", False),
    ("FM05 Insertion accessory (guidewire, stylet, sheath, needle, tunneller, tip-location system)", False),
    ("FM06 Malposition / migration", False),
    ("FM07 Difficult insertion or removal, no defect found", False),
    ("FM08 Packaging, labelling or assembly defect found before use", False),
    ("FM09 Infection   FM10 Thrombosis   FM11 Vessel injury / extravasation", False),
    ("FM12 Other / insufficient information   FM13 Tissue / hypersensitivity reaction", False),
    ("", False),
    ("Phase: P1 before use, P2 insertion, P3 in use, P4 removal, P5 unknown", False),
    ("Severity: S1 negligible, S2 minor, S3 serious (medical intervention), S4 critical (life-threatening or invasive retrieval), S5 death", False),
    ("Contributor: DEV device, USE use error, PAT patient, UNK unknown", False),
    ("Retained fragment / user harm: Y or N", False),
    ("Rules: breach beats deformation and occlusion; device failure beats clinical complication; cause beats effect;", False),
    ("leak at insertion site with no visible breach = FM01 low confidence; leak with no location = FM12 low confidence.", False),
    ("Repair or trim of the catheter = at least S2.", False),
]
for i, (t, b) in enumerate(lines, 1):
    ins.cell(i, 1, t).font = Font(bold=b, size=12 if i == 1 else 11)
ins.column_dimensions["A"].width = 120

ws = wb.create_sheet("Coding")
head = ["#", "Report number", "Event type", "Device", "FDA device problem codes", "FDA patient problem codes",
        "Event description", "Manufacturer narrative (summary)",
        "Primary FM", "Secondary FM", "Phase", "Severity", "Contributor", "Retained fragment", "User harm",
        "Confidence", "My notes"]
ws.append(head)
for i, k in enumerate(pick, 1):
    r = rep[k]
    ws.append([i, k, r["event_type"], r["brand_name"], r["device_problems"], r["patient_problems"],
               dedupe(r["event_description"])[:3000], mfr_summary(r["manufacturer_narrative"])[:1200]] + [""] * 9)
fill = PatternFill("solid", fgColor="FFF2CC")
for c in ws[1]:
    c.font = Font(bold=True)
widths = [5, 22, 12, 28, 26, 26, 70, 45, 11, 11, 8, 9, 11, 10, 9, 11, 30]
for col, w in zip("ABCDEFGHIJKLMNOPQ", widths):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in row[8:]:
        c.fill = fill
ws.freeze_panes = "C2"
fms = ",".join(f"FM{n:02d}" for n in range(1, 14))
dv = [(f'"{fms}"', "I"), (f'"none,{fms}"', "J"), ('"P1,P2,P3,P4,P5"', "K"), ('"S1,S2,S3,S4,S5"', "L"),
      ('"DEV,USE,PAT,UNK"', "M"), ('"Y,N"', "N"), ('"Y,N"', "O"), ('"high,medium,low"', "P")]
for formula, col in dv:
    v = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(v)
    v.add(f"{col}2:{col}101")
wb.save(ROOT / "data" / "step5_manual_coding.xlsx")
print("Wrote", len(pick), "reports")
