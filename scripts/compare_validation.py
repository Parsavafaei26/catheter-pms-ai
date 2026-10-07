"""Step 5: compare blind manual coding with AI coding.

Usage: python3 scripts/compare_validation.py [ai_labels.csv] [output.csv]

Reads data/step5_manual_coding_completed.xlsx (manual) and data/processed/ai_labels_full.csv (AI),
prints per-field agreement and Cohen's kappa, and writes data/processed/validation_comparison.csv.
"""
import csv, sys
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
FIELDS = [("primary_fm", 9), ("secondary_fm", 10), ("phase", 11), ("severity", 12),
          ("contributor", 13), ("retained_fragment", 14), ("user_harm", 15)]


def kappa(a, b):
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / n ** 2
    return (po - pe) / (1 - pe) if pe < 1 else 1.0


AI_FILE = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data/processed/ai_labels_full.csv"
OUT_FILE = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "data/processed/validation_comparison.csv"
ai = {r["report_number"]: r for r in csv.DictReader(open(AI_FILE, encoding="utf-8"))}
ws = load_workbook(ROOT / "data/step5_manual_coding_completed.xlsx")["Coding"]
rows = []
for r in range(2, ws.max_row + 1):
    rid = ws.cell(r, 2).value
    if not rid:
        continue
    man = {f: (str(ws.cell(r, c).value or "").strip()) for f, c in FIELDS}
    man["note"] = ws.cell(r, 17).value or ""
    rows.append((rid, man, ai[rid]))

print(f"Reports compared: {len(rows)}\n")
print(f"{'Field':20}{'Agreement':>10}{'Kappa':>8}")
summary = {}
for f, _ in FIELDS:
    a = [m[f] for _, m, _ in rows]
    b = [x[f] for _, _, x in rows]
    ag = sum(x == y for x, y in zip(a, b)) / len(rows)
    k = kappa(a, b)
    summary[f] = (ag, k)
    print(f"{f:20}{ag:>9.0%}{k:>8.2f}")
sev = lambda s: int(s[1])
w1 = sum(abs(sev(m['severity']) - sev(x['severity'])) <= 1 for _, m, x in rows) / len(rows)
print(f"{'severity within 1':20}{w1:>9.0%}")
diff = Counter(sev(m['severity']) - sev(x['severity']) for _, m, x in rows)
print("severity difference (manual - AI):", dict(sorted(diff.items())))
allf = sum(all(m[f] == x[f] for f, _ in FIELDS) for _, m, x in rows) / len(rows)
print(f"all 7 fields identical: {allf:.0%}")

with open(OUT_FILE, "w", newline="", encoding="utf-8") as fh:
    w = csv.writer(fh)
    w.writerow(["report_number"] + [f"{f}_{s}" for f, _ in FIELDS for s in ("manual", "ai")] + ["manual_note", "ai_note"])
    for rid, m, x in rows:
        w.writerow([rid] + [v for f, _ in FIELDS for v in (m[f], x[f])] + [m["note"], x["note"]])

print("\nPrimary failure mode disagreements:")
for rid, m, x in rows:
    if m["primary_fm"] != x["primary_fm"]:
        print(f"  {rid}: manual {m['primary_fm']} / AI {x['primary_fm']} | {m['note'][:90]}")
print("\nPrimary FM confusion (manual -> AI):", Counter((m['primary_fm'], x['primary_fm']) for _, m, x in rows if m['primary_fm'] != x['primary_fm']))
